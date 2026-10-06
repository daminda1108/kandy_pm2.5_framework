"""Apply the departmental format pandoc cannot express, then measure it in Word.

Two stages.

STRUCTURE (python-docx, on pandoc's output). Acts on the marker paragraphs assemble.py leaves:

    [[SECTION-FRONT]]   ends the title-page section: the title page carries no page number
    [[SECTION-BODY]]    ends the front-matter section, numbered in lower-case roman from i;
                        everything after it is numbered in Arabic from 1, starting at the
                        Table of Contents as the guidelines require
    [[TOC]] [[LOF]] [[LOT]]   replaced by real Word fields, so the lists carry page numbers
                        and update like any Word document

Every figure and table caption also gains a hidden TC field holding its short caption, which is
what the List of Figures and List of Tables are built from. Building those lists from the whole
caption would reproduce paragraphs of caption text in the front matter.

FINALISE (Word, through COM; WINWORD.EXE is on this machine). Opens the document, sets the
contents styles, updates every field twice (the contents change length, which moves the pages
they point at), saves, exports the PDF, and MEASURES the result against the guidelines:

    body 40 to 75 pages, exclusive of appendices     ← Arabic pages before Appendix A
    Introduction not more than 40 % of the pages
    Results and Discussion not less than 30 %
    abstract on a single page

The page count is read from the document, not estimated from words. An estimate is how this
thesis reached 136 pages against a 75-page limit without anyone noticing.

Usage: python postprocess.py --thesis a      (build_docx.py calls it)
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).resolve().parent

PAGES_MIN, PAGES_MAX = 40, 75
INTRO_MAX, RESULTS_MIN = 0.40, 0.30
CAPTION = re.compile(r"^(Figure|Table)\s+([0-9A-Z]+\.\d+)\.")


# ── fields ────────────────────────────────────────────────────────────────────────────────

def _run(parent, child):
    r = OxmlElement("w:r")
    r.append(child)
    parent.append(r)
    return r


def _fldchar(kind: str, dirty: bool = False):
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), kind)
    if dirty:
        fc.set(qn("w:dirty"), "true")
    return fc


def _instr(code: str):
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {code} "
    return it


def add_field(p_el, code: str, placeholder: str = "") -> None:
    """Append a complete field (begin, code, separate, result, end) to a paragraph element."""
    _run(p_el, _fldchar("begin", dirty=True))
    _run(p_el, _instr(code))
    _run(p_el, _fldchar("separate"))
    t = OxmlElement("w:t")
    t.text = placeholder
    _run(p_el, t)
    _run(p_el, _fldchar("end"))


def _clear_runs(p_el) -> None:
    for child in list(p_el):
        if child.tag != qn("w:pPr"):
            p_el.remove(child)


def _set_style(p, name: str, doc) -> None:
    p.style = doc.styles[name]


# ── sections and page numbers ─────────────────────────────────────────────────────────────

def _section_break(p_el, body_sectpr, fmt: str | None) -> None:
    """Make this paragraph the last one of a section, with the page-number format given."""
    sp = copy.deepcopy(body_sectpr)
    for ref in sp.findall(qn("w:headerReference")) + sp.findall(qn("w:footerReference")):
        sp.remove(ref)
    for old in sp.findall(qn("w:pgNumType")):
        sp.remove(old)
    for old in sp.findall(qn("w:type")):
        sp.remove(old)
    typ = OxmlElement("w:type")
    typ.set(qn("w:val"), "nextPage")
    sp.insert(0, typ)
    if fmt:
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:fmt"), fmt)
        pg.set(qn("w:start"), "1")
        sp.append(pg)
    ppr = p_el.get_or_add_pPr()
    ppr.append(sp)


def _set_pgnum(sectpr, fmt: str) -> None:
    for old in sectpr.findall(qn("w:pgNumType")):
        sectpr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), "1")
    sectpr.append(pg)


def _footer_page_number(section) -> None:
    section.footer.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]
    _clear_runs(fp._p)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp._p, "PAGE", "1")
    for r in fp.runs:
        r.font.name, r.font.size = "Times New Roman", Pt(12)


def structure(src: Path, dst: Path, meta: dict) -> list[str]:
    d = docx.Document(str(src))
    body = d.element.body
    body_sectpr = body.find(qn("w:sectPr"))
    problems: list[str] = []

    markers = {}
    for p in d.paragraphs:
        m = re.fullmatch(r"\s*\[\[([A-Z\-]+)\]\]\s*", p.text)
        if m:
            markers.setdefault(m.group(1), []).append(p)
    for need in ("SECTION-FRONT", "SECTION-BODY", "TOC", "LOF", "LOT"):
        if len(markers.get(need, [])) != 1:
            problems.append(f"marker [[{need}]] found {len(markers.get(need, []))} times, "
                            f"expected once")
    if problems:
        return problems

    # Contents fields. \o "1-2": chapters and their sections. Listing the third level as well
    # ran the contents to four pages, most of it appendix subsections, inside a body that has a
    # page limit. \f F and \f T read the TC fields written into the captions below.
    for tag, code in (("TOC", 'TOC \\o "1-2" \\h \\z \\u'),
                      ("LOF", "TOC \\h \\z \\f F"),
                      ("LOT", "TOC \\h \\z \\f T")):
        p = markers[tag][0]
        _clear_runs(p._p)
        _set_style(p, "Normal", d)
        add_field(p._p, code, "Contents are generated when the document is updated.")

    # TC fields on captions, one per figure and table, carrying the short caption.
    short = {c["label"]: c["short"] for c in meta["captions"]}
    found = set()
    for p in d.paragraphs:
        m = CAPTION.match(p.text)
        if not m or p.style.name not in ("Image Caption", "Table Caption", "Caption"):
            continue
        label = f"{m.group(1)} {m.group(2)}"
        if label not in short or label in found:
            continue
        found.add(label)
        flag = "F" if m.group(1) == "Figure" else "T"
        entry = f"{label}. {short[label]}".replace('"', "'")
        add_field(p._p, f'TC "{entry}" \\f {flag} \\l 1', "")
    missing = sorted(set(short) - found)
    if missing:
        problems.append(f"{len(missing)} caption(s) not found in the document: "
                        + ", ".join(missing[:8]))

    # Sections. The marker paragraph is emptied and carries the break that ENDS its section.
    for tag, fmt in (("SECTION-FRONT", None), ("SECTION-BODY", "lowerRoman")):
        p = markers[tag][0]
        _clear_runs(p._p)
        _section_break(p._p, body_sectpr, fmt)
    _set_pgnum(body_sectpr, "decimal")

    secs = d.sections
    if len(secs) != 3:
        problems.append(f"expected 3 sections, found {len(secs)}")
        return problems
    secs[0].footer.is_linked_to_previous = False           # title page: no number
    for p in secs[0].footer.paragraphs:
        _clear_runs(p._p)
    _footer_page_number(secs[1])                            # roman
    _footer_page_number(secs[2])                            # Arabic

    d.save(str(dst))
    return problems


# ── Word ──────────────────────────────────────────────────────────────────────────────────

WD_STYLE_TOC = (-20, -21, -22)          # wdStyleTOC1..3
WD_STYLE_TABLE_OF_FIGURES = -36
WD_PAGE_PHYSICAL = 3                     # wdActiveEndPageNumber
WD_PAGE_SHOWN = 1                        # wdActiveEndAdjustedPageNumber


def finalise(docx_path: Path, pdf_path: Path, meta: dict) -> dict:
    import pythoncom
    import win32com.client as win32

    pythoncom.CoInitialize()
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    report: dict = {}
    try:
        doc = word.Documents.Open(str(docx_path.resolve()), False, False)
        for sid in WD_STYLE_TOC + (WD_STYLE_TABLE_OF_FIGURES,):
            st = doc.Styles(sid)
            st.Font.Name = "Times New Roman"
            st.Font.Size = 12
            st.ParagraphFormat.LineSpacingRule = 0          # single
            st.ParagraphFormat.SpaceAfter = 4
            st.ParagraphFormat.Alignment = 0                # left
        # Not bold: the List of Figures and List of Tables are built from level-1 TC entries and
        # so share the TOC 1 style; bold made every figure entry bold. Chapter titles are already
        # set apart in the contents by being upper case.
        doc.Styles(WD_STYLE_TOC[0]).Font.Bold = False
        for _ in range(2):
            doc.Fields.Update()
            for i in range(1, doc.TablesOfContents.Count + 1):
                doc.TablesOfContents(i).Update()
            doc.Repaginate()

        # Chapter starts, from the contents list: its page numbers are the ones a reader sees,
        # and the Arabic run starts at the contents page, so differences between them are
        # physical page counts.
        toc = doc.TablesOfContents(1).Range.Text
        starts = []
        for line in toc.split("\r"):
            m = re.match(r"^(.*?)\t(\d+)\s*$", line.strip())
            if m:
                starts.append((m.group(1).strip(), int(m.group(2))))
        total_phys = doc.ComputeStatistics(2)
        last_shown = doc.Content.Characters.Last.Information(WD_PAGE_SHOWN)

        # abstract: every Abstract Text paragraph on one physical page
        abs_pages = set()
        for i in range(1, min(doc.Paragraphs.Count, 400) + 1):
            p = doc.Paragraphs(i)
            if p.Style.NameLocal == "Abstract Text":
                abs_pages.add(p.Range.Information(WD_PAGE_PHYSICAL))

        doc.Save()
        doc.ExportAsFixedFormat(str(pdf_path.resolve()), 17)   # wdExportFormatPDF
        doc.Close(False)
    finally:
        word.Quit()
        pythoncom.CoUninitialize()

    chap = [(t, pg) for t, pg in starts if re.match(r"^\d+\.\s", t)]
    apps = [(t, pg) for t, pg in starts if t.startswith("APPENDIX")]
    body_pages = (apps[0][1] - 1) if apps else last_shown

    def span(word_):
        for i, (t, pg) in enumerate(chap):
            if word_ in t:
                nxt = chap[i + 1][1] if i + 1 < len(chap) else next(
                    (pg2 for t2, pg2 in starts if pg2 > pg and not re.match(r"^\d+\.\s", t2)),
                    body_pages + 1)
                return nxt - pg
        return None

    intro, results = span("INTRODUCTION"), span("RESULTS AND DISCUSSION")
    report = {
        "physical_pages_total": total_phys,
        "arabic_pages_total": last_shown,
        "body_pages_excl_appendices": body_pages,
        "chapter_starts": chap,
        "appendix_starts": apps,
        "introduction_pages": intro,
        "results_pages": results,
        "introduction_share": round(intro / body_pages, 3) if intro else None,
        "results_share": round(results / body_pages, 3) if results else None,
        "abstract_physical_pages": sorted(abs_pages),
    }
    fails = []
    if not PAGES_MIN <= body_pages <= PAGES_MAX:
        fails.append(f"body is {body_pages} pages excluding appendices; "
                     f"required {PAGES_MIN} to {PAGES_MAX}")
    if intro is None or report["introduction_share"] > INTRO_MAX:
        fails.append(f"Introduction is {intro} pages ({report['introduction_share']}); "
                     f"limit {INTRO_MAX:.0%}")
    if results is None or report["results_share"] < RESULTS_MIN:
        fails.append(f"Results and Discussion is {results} pages ({report['results_share']}); "
                     f"floor {RESULTS_MIN:.0%}")
    if len(abs_pages) != 1:
        fails.append(f"abstract spans pages {sorted(abs_pages)}; it must fit on one")
    report["fails"] = fails
    return report


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--thesis", required=True)
    ap.add_argument("--no-word", action="store_true", help="structure only; skip Word")
    a = ap.parse_args()
    out = HERE / a.thesis
    meta = json.loads((out / "meta.json").read_text(encoding="utf-8"))
    src = out / "pandoc.docx"
    dst = out / f"thesis_{a.thesis}.docx"
    pdf = out / f"thesis_{a.thesis}.pdf"

    problems = structure(src, dst, meta)
    if problems:
        print("STRUCTURE FAILED:")
        for p in problems:
            print(f"  {p}")
        return 1
    print(f"  structured  {dst.relative_to(HERE.parent)}")
    if a.no_word:
        return 0

    # A PDF held open by a viewer cannot be overwritten, and Word then fails the whole export
    # (2026-09-19). Write beside it instead and say so, so the compliance check still runs.
    if pdf.exists():
        try:
            with open(pdf, "r+b"):
                pass
        except OSError:
            pdf = pdf.with_name(f"{pdf.stem}_new.pdf")
            print(f"  NOTE        thesis_{a.thesis}.pdf is open elsewhere; writing {pdf.name}")
    rep = finalise(dst, pdf, meta)
    (out / "compliance.json").write_text(json.dumps(rep, indent=1), encoding="utf-8")
    print(f"  pages       {rep['physical_pages_total']} physical; "
          f"{rep['body_pages_excl_appendices']} Arabic before the appendices "
          f"(required {PAGES_MIN}-{PAGES_MAX})")
    print(f"  intro       {rep['introduction_pages']} pages, share {rep['introduction_share']} "
          f"(max {INTRO_MAX})")
    print(f"  results     {rep['results_pages']} pages, share {rep['results_share']} "
          f"(min {RESULTS_MIN})")
    print(f"  abstract    on page(s) {rep['abstract_physical_pages']}")
    print(f"  wrote       {pdf.relative_to(HERE.parent)}")
    if rep["fails"]:
        print("\nNOT COMPLIANT:")
        for f in rep["fails"]:
            print(f"  {f}")
        return 2
    print("  COMPLIANT with the page, proportion and abstract rules")
    return 0


if __name__ == "__main__":
    sys.exit(main())
